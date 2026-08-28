from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

def export_database(database, filename):

    workbook = Workbook()

    # Gunakan sheet bawaan sebagai sheet Slalom
    workbook.active.title = "Slalom"

    export_slalom(workbook, database)
    export_endurance(workbook, database)
    export_brake(workbook, database)
    export_hill(workbook, database)

    workbook.save(filename)


def export_slalom(workbook, database):

    sheet = workbook["Slalom"]

    headers = [
        "No Urut",
        "Nama Team",
        "Universitas",
        "Waktu (s)",
        "Speed (km/h)",
        "Status"
    ]

    sheet.append(headers)

    # Header Bold
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    rows = database.get_slalom()

    for row in rows:
        sheet.append(row)
    autofit_columns(sheet)


def export_endurance(workbook, database):

    sheet = workbook.create_sheet("Endurance")

    headers = [
        "No Urut",
        "Nama Team",
        "Universitas",
        "Lap 1 (s)",
        "Lap 2 (s)",
        "Total (s)",
        "Status"
    ]

    sheet.append(headers)

    # Header Bold
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    rows = database.get_lap()

    for row in rows:
        sheet.append(row)
    autofit_columns(sheet)


def export_brake(workbook, database):

    sheet = workbook.create_sheet("Akselerasi Deselerasi")

    headers = [
        "No Urut",
        "Nama Team",
        "Universitas",
        "Acceleration (m/s²)",
        "Deceleration (m/s²)",
        "Status"
    ]

    sheet.append(headers)

    # Header Bold
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    rows = database.get_brake()

    for row in rows:
        sheet.append(row)
    autofit_columns(sheet)


def export_hill(workbook, database):

    sheet = workbook.create_sheet("Daya Tanjak")

    headers = [
        "No Urut",
        "Nama Team",
        "Universitas",
        "Massa (kg)",
        "Waktu (s)",
        "Kecepatan (km/h)",
        "Daya (kW)",
        "Status"
    ]

    sheet.append(headers)

    # Header Bold
    for cell in sheet[1]:
        cell.font = Font(bold=True)

    rows = database.get_hill()

    for row in rows:
        sheet.append(row)
    autofit_columns(sheet)

def autofit_columns(sheet):

    for column_cells in sheet.columns:

        max_length = 0

        column = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            value = cell.value

            if value is None:
                continue

            length = len(str(value))

            if length > max_length:
                max_length = length

        sheet.column_dimensions[column].width = max_length + 2
